#!/usr/bin/env python
"""Build the AMDL-TPNE paper as a DOCX in the reference PDF's IEEE two-column format."""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Emu

import content as C
from common import (parse_runs, normalise, FONT, PAGE_W, PAGE_H, MARGIN_L, MARGIN_R,
                    MARGIN_T, MARGIN_B, COL_GAP, COL_W, BODY_SIZE, BODY_LEAD, TITLE_SIZE,
                    AUTHOR_SIZE, AFFIL_SIZE, CORR_SIZE, HEADING_SIZE, SUBSEC_SIZE,
                    ABSTRACT_SIZE, TABLE_SIZE, CAPTION_SIZE, REF_SIZE)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "AMDL-TPNE_Paper.docx")


# ------------------------------------------------------------------ low level
def set_font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def set_columns(section, num, gap_pt=COL_GAP):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:cols")):
        sectPr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(gap_pt * 20)))
    if num > 1:
        cols.set(qn("w:equalWidth"), "1")
        cols.set(qn("w:sep"), "0")
    sectPr.append(cols)


def para(doc, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         line=BODY_LEAD, indent_first=0.0, indent_left=0.0, keep_next=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = Pt(line)
    fmt.alignment = align
    fmt.first_line_indent = Pt(indent_first)
    fmt.left_indent = Pt(indent_left)
    fmt.keep_with_next = keep_next
    return p


def add_runs(p, text, size, bold=False, italic=False, sc=False):
    for r in parse_runs(normalise(text), bold=bold, italic=italic, size=size, sc=sc):
        if not r["text"]:
            continue
        for i, piece in enumerate(r["text"].split("\n")):
            if i:
                br = p.add_run()
                set_font(br)
                br.font.size = Pt(r["size"])
                br.add_break()
            if not piece:
                continue
            run = p.add_run(piece)
            set_font(run)
            run.font.size = Pt(r["size"])
            run.bold = r["bold"]
            run.italic = r["italic"]
            if r["sc"]:
                run.font.small_caps = True
            if r["sup"]:
                run.font.superscript = True
            if r["sub"]:
                run.font.subscript = True


def para_runs(p, text, size, bold=False, italic=False, sc=False):
    add_runs(p, text, size, bold, italic, sc)


def horizontal_rule(paragraph, position, sz=8, color="000000"):
    """Add a top/bottom border to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    borders = pPr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        pPr.append(borders)
    el = OxmlElement("w:" + position)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    borders.append(el)


def cell_border(cell, edge, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tblPr.append(borders)


def add_page_field(paragraph, size=7.5):
    """Insert a 'Page N' field, centred in the footer."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_font(run)
    run.font.size = Pt(size)

    r1 = paragraph.add_run()
    set_font(r1)
    r1.font.size = Pt(size)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r1._element.addnext(fld)

    run2 = paragraph.add_run(" of ")
    set_font(run2)
    run2.font.size = Pt(size)
    r2 = paragraph.add_run()
    set_font(r2)
    r2.font.size = Pt(size)
    fld2 = OxmlElement("w:fldSimple")
    fld2.set(qn("w:instr"), "NUMPAGES")
    r2._element.addnext(fld2)


def style_footnotes(doc):
    try:
        st = doc.styles["Footnote Text"]
        st.font.size = Pt(CAPTION_SIZE)
        st.font.name = FONT
        st.paragraph_format.line_spacing = Pt(CAPTION_SIZE + 1.5)
    except KeyError:
        pass


# ------------------------------------------------------------------ blocks
def build_title_block(doc):
    p = para(doc, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER, line=18)
    add_runs(p, C.TITLE, TITLE_SIZE, bold=True)

    p = para(doc, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER, line=13)
    for i, (name, mark, corr) in enumerate(C.AUTHORS):
        if i:
            add_runs(p, ",  ", AUTHOR_SIZE)
        star = "<sup>*</sup>" if corr else ""
        add_runs(p, f"{name}<sup>{mark}{star}</sup>", AUTHOR_SIZE)

    for mark, aff in C.AFFILIATIONS:
        p = para(doc, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER, line=10.5)
        add_runs(p, f"<sup>{mark}</sup>{aff}", AFFIL_SIZE)

    p = para(doc, space_before=4, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER, line=10.5)
    add_runs(p, f"Corresponding author: {C.CORRESPONDING['name']}", CORR_SIZE, italic=True)
    p = para(doc, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER, line=10.5)
    add_runs(p, f"Email: {C.CORRESPONDING['email']}"
                f"&nbsp;&nbsp;&nbsp;Tel: {C.CORRESPONDING['phone']}", CORR_SIZE, italic=True)

    # Abstract
    p = para(doc, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line=BODY_LEAD, indent_left=10, indent_first=0)
    add_runs(p, "<b>ABSTRACT</b>&nbsp;&mdash;&nbsp;", ABSTRACT_SIZE, sc=True)
    add_runs(p, C.ABSTRACT, ABSTRACT_SIZE)
    p.paragraph_format.right_indent = Pt(10)

    p = para(doc, space_after=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line=BODY_LEAD, indent_left=10)
    add_runs(p, "<b>KEYWORDS</b>&nbsp;&mdash;&nbsp;", ABSTRACT_SIZE, sc=True)
    add_runs(p, C.KEYWORDS, ABSTRACT_SIZE, italic=True)
    p.paragraph_format.right_indent = Pt(10)


def build_heading(doc, text, level):
    size = HEADING_SIZE if level == 1 else SUBSEC_SIZE
    before = 9 if level == 1 else 6
    p = para(doc, space_before=before, space_after=4,
             align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT,
             line=size + 2.5, keep_next=True)
    add_runs(p, text, size, bold=True, sc=(level == 1))


def build_figure(doc, path, number, caption):
    full = os.path.join(HERE, path)
    p = para(doc, space_before=8, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER, line=10)
    p.add_run().add_picture(full, width=Inches(COL_W / 72.0))
    p.paragraph_format.keep_with_next = True

    cp = para(doc, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=9.4)
    add_runs(cp, caption, CAPTION_SIZE)


def build_table(doc, spec):
    cap = spec["caption"]
    size = spec.get("size", TABLE_SIZE)
    widths = spec.get("widths")
    head, rows = spec["head"], spec["rows"]

    cp = para(doc, space_before=9, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER,
              line=9.4, keep_next=True)
    add_runs(cp, cap, CAPTION_SIZE)

    ncols = len(head)
    table = doc.add_table(rows=0, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_borders(table)

    if widths is None:
        widths = [1.0 / ncols] * ncols
    total = sum(widths)
    col_pts = [COL_W * w / total for w in widths]

    def fill_row(cells, values, bold=False):
        for cell, val, w in zip(cells, values, col_pts):
            cell.width = Pt(w)
            cell.paragraphs[0].text = ""
            pp = cell.paragraphs[0]
            pf = pp.paragraph_format
            pf.space_before = Pt(1.2)
            pf.space_after = Pt(1.2)
            pf.line_spacing = Pt(size + 2.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(pp, val, size, bold=bold)

    hr = table.add_row()
    fill_row(hr.cells, head, bold=True)
    for cell in hr.cells:
        cell_border(cell, "top", sz=8)
        cell_border(cell, "bottom", sz=6)

    for i, row in enumerate(rows):
        r = table.add_row()
        fill_row(r.cells, row)
        if i == len(rows) - 1:
            for cell in r.cells:
                cell_border(cell, "bottom", sz=8)

    para(doc, space_before=0, space_after=8, line=6).add_run()


def build_body(doc):
    for kind, payload in C.BODY:
        if kind == "h":
            level, text = payload
            build_heading(doc, text, level)
        elif kind == "p":
            p = para(doc, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=BODY_LEAD)
            add_runs(p, payload, BODY_SIZE)
        elif kind == "eq":
            p = para(doc, space_before=4, space_after=4,
                     align=WD_ALIGN_PARAGRAPH.CENTER, line=BODY_SIZE + 2.6)
            add_runs(p, payload, BODY_SIZE)
        elif kind == "fig":
            path, number, caption = payload
            build_figure(doc, path, number, caption)
        elif kind == "tab":
            build_table(doc, payload)


def build_references(doc):
    p = para(doc, space_before=9, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER,
             line=HEADING_SIZE + 2.5, keep_next=True)
    add_runs(p, "REFERENCES", HEADING_SIZE, bold=True, sc=True)

    for i, ref in enumerate(C.REFERENCES, 1):
        p = para(doc, space_after=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 line=REF_SIZE + 2.0, indent_left=14, indent_first=-14)
        add_runs(p, f"[{i}]\t{normalise(ref)}", REF_SIZE)


# ------------------------------------------------------------------ document
def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_SIZE)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(BODY_LEAD)

    s0 = doc.sections[0]
    s0.page_width = Pt(PAGE_W)
    s0.page_height = Pt(PAGE_H)
    s0.left_margin = Pt(MARGIN_L)
    s0.right_margin = Pt(MARGIN_R)
    s0.top_margin = Pt(MARGIN_T)
    s0.bottom_margin = Pt(MARGIN_B)
    set_columns(s0, 1)

    build_title_block(doc)

    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    s1.page_width = Pt(PAGE_W)
    s1.page_height = Pt(PAGE_H)
    s1.left_margin = Pt(MARGIN_L)
    s1.right_margin = Pt(MARGIN_R)
    s1.top_margin = Pt(MARGIN_T)
    s1.bottom_margin = Pt(MARGIN_B)
    set_columns(s1, 2)

    build_body(doc)
    build_references(doc)

    for sec in doc.sections:
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        for r in list(fp.runs):
            r._element.getparent().remove(r._element)
        add_page_field(fp)

    style_footnotes(doc)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()